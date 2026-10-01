import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

document = Document(r"D:\Browser Download\Đề cương đồ án SE405.docx")
for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text:
        print(text)
for table_index, table in enumerate(document.tables, start=1):
    print(f"\n[TABLE {table_index}]")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
